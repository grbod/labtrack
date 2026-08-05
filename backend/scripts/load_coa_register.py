#!/usr/bin/env python3
"""Wipe per-lot lab data and load the real COA Register spreadsheet.

Reads '2.4.4 New COA Register.xlsx', groups rows into standard lots, parent lots
(shared Lot number -> sublots = RefIDs) and multi-SKU composites (shared C-ref),
auto-creates missing products (with their real test panel), and recreates lots /
sublots / test results / releases with the correct status.

Status rules (per group, from the lead row + the ToPrint flag):
  * QC Approval present AND not "NEEDS METALS"      -> RELEASED  (+ COARelease/product)
  * ToPrint == "NEEDS METALS"                       -> PARTIAL_RESULTS (micro in, metals pending)
  * no QC, but the row already has results          -> AWAITING_RELEASE (shows in Release Queue)
  * no QC and no results                            -> AWAITING_RESULTS (shows in Sample Tracker)

Usage:
    cd backend
    .venv/bin/python scripts/load_coa_register.py --dry-run     # report only (default)
    .venv/bin/python scripts/load_coa_register.py --commit      # backup + wipe + load
    .venv/bin/python scripts/load_coa_register.py --append-from-row 1896 --commit
"""

from __future__ import annotations

import argparse
import calendar
import csv
import os
import sqlite3
import sys
from collections import Counter, OrderedDict, defaultdict
from datetime import date, datetime
import re

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl  # noqa: E402
from sqlalchemy import text  # noqa: E402

from app.database import SessionLocal  # noqa: E402
from app.models.coa_release import COARelease  # noqa: E402
from app.models.enums import (  # noqa: E402
    COAReleaseStatus,
    LotStatus,
    LotType,
    TestResultStatus,
)
from app.models.lab_test_type import LabTestType  # noqa: E402
from app.models.lot import Lot, LotProduct, Sublot  # noqa: E402
from app.models.product import Product  # noqa: E402
from app.models.product_test_spec import ProductTestSpecification  # noqa: E402
from app.models.test_result import TestResult  # noqa: E402
from app.models.user import User  # noqa: E402

DEFAULT_XLSX = "/Users/gregsimek/Downloads/2.4.4 New COA Register.xlsx"
BACKEND_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DB_PATH = os.path.join(BACKEND_DIR, "labtrack.db")
BACKUP_DIR = os.path.join(BACKEND_DIR, "backups")

C_TOPRINT, C_REFID, C_BRAND, C_PRODUCT, C_FLAVOR, C_SIZE, C_LOT = 0, 1, 2, 3, 4, 5, 6
C_MFG, C_EXP, C_RELEASE, C_QC = 7, 8, 9, 14

TEST_COLS = [
    (10, "Total Plate Count"),
    (11, "Yeast & Mold"),
    (12, "Escherichia coli"),
    (13, "Salmonella spp."),
    (15, "Gluten"),
    (16, "Staphylococcus aureus"),
    (17, "Total Coliform Count"),
    (18, "Arsenic"),
    (19, "Cadmium"),
    (20, "Lead"),
    (21, "Mercury"),
    (23, "Caffeine"),
]
# Panel every product is expected to be tested for when nothing else is known.
MICRO_COMMON = [
    "Total Plate Count",
    "Yeast & Mold",
    "Escherichia coli",
    "Salmonella spp.",
]
METALS = ["Arsenic", "Cadmium", "Lead", "Mercury"]

WIPE_ORDER = [
    "email_history",
    "release_sensory_attests",
    "coa_snapshots",
    "retest_items",
    "retest_requests",
    "coa_releases",
    "coa_history",
    "test_results",
    "sublots",
    "lot_products",
    "audit_annotations",
    "audit_logs",
    "parsing_queue",
    "lots",
    "daane_coc_daily_counters",
    "coa_serial_counters",
]


# --------------------------------------------------------------------------- #
class RegisterRow(list):
    """A spreadsheet row that still behaves like the original list of cells."""

    def __init__(self, values, excel_row):
        super().__init__(values)
        self.excel_row = excel_row


def row_ref(row) -> str:
    return str(getattr(row, "excel_row", "?"))


def cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def to_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        s = v.strip()
        if not s:
            return None
        for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y"):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                pass
        # The register sometimes contains impossible month-end dates such as
        # 06/31/2029, and one duplicated slash typo. Preserve the intended
        # month/year and clamp only an invalid day to that month's final day.
        match = re.fullmatch(r"(\d{1,2})/+(\d{1,2})/+(\d{2,4})", s)
        if match:
            month_text, day_text, year_text = match.groups()
            month, day, year = int(month_text), int(day_text), int(year_text)
            if year < 100:
                year += 2000
            elif len(year_text) == 3 and 200 <= year <= 209:
                # A recurring register typo drops the third digit from a
                # 2020s year, e.g. 07/20/206 means 2026.
                year += 1820
            if 1 <= month <= 12 and day >= 1:
                return date(year, month, min(day, calendar.monthrange(year, month)[1]))
    return None


def is_blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def is_composite_ref(refid) -> bool:
    return str(refid).strip().upper().startswith("C")


def display_name(brand, product, flavor, size) -> str:
    return " - ".join(str(p).strip() for p in (brand, product, flavor, size) if p)


def sku_key(row):
    def n(i):
        return str(row[i]).strip().lower() if not is_blank(row[i]) else ""

    return (n(C_BRAND), n(C_PRODUCT), n(C_FLAVOR), n(C_SIZE))


def load_rows(xlsx_path, append_from_row=None):
    if str(xlsx_path).lower().endswith(".csv"):
        rows = []
        with open(xlsx_path, newline="", encoding="utf-8-sig") as fh:
            reader = csv.reader(fh)
            next(reader, None)
            for csv_row, r in enumerate(reader, start=2):
                if append_from_row is not None and csv_row < append_from_row:
                    continue
                if any(not is_blank(v) for v in r):
                    rows.append(RegisterRow(r, csv_row))
        return rows

    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb["Sheet1"] if "Sheet1" in wb.sheetnames else wb.active
    rows = []
    for excel_row, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if append_from_row is not None and excel_row < append_from_row:
            continue
        if any(not is_blank(v) for v in r):
            rows.append(RegisterRow(list(r), excel_row))
    return rows


def group_rows(rows):
    flags = []
    c_groups, shared_ref_groups, rest = OrderedDict(), OrderedDict(), []
    ref_occurrences = defaultdict(list)
    for r in rows:
        ref = cell_str(r[C_REFID])
        if ref and not is_composite_ref(ref):
            ref_occurrences[ref].append(r)

    shared_parent_refs = set()
    for ref, ref_rows in ref_occurrences.items():
        excel_rows = [getattr(r, "excel_row", 0) for r in ref_rows]
        is_contiguous = excel_rows == list(range(excel_rows[0], excel_rows[-1] + 1))
        if (
            len(ref_rows) > 1
            and is_contiguous
            and len({sku_key(r) for r in ref_rows}) == 1
            and len({cell_str(r[C_LOT]) for r in ref_rows}) > 1
        ):
            shared_parent_refs.add(ref)

    for r in rows:
        ref = cell_str(r[C_REFID])
        if is_composite_ref(ref):
            c_groups.setdefault(ref, []).append(r)
        elif ref in shared_parent_refs:
            shared_ref_groups.setdefault(ref, []).append(r)
        else:
            rest.append(r)
    groups = []
    for cref, grp in c_groups.items():
        if len({sku_key(x) for x in grp}) > 1:
            groups.append({"kind": "composite", "key": cref, "rows": grp})
        else:
            groups.append(
                {
                    "kind": "parent",
                    "key": cref,
                    "rows": grp,
                    "ref_override": cref,
                    "lotnum_override": cref,
                }
            )
            flags.append(
                f"same-SKU composite {cref} loaded as a PARENT lot ({len(grp)} batches)"
            )
    for ref, grp in shared_ref_groups.items():
        groups.append(
            {
                "kind": "parent",
                "key": ref,
                "rows": grp,
                "ref_override": ref,
                "lotnum_override": ref,
            }
        )
        flags.append(
            f"shared non-C RefID {ref} loaded as a PARENT lot ({len(grp)} batches)"
        )
    current_key, current_lot, current_rows = None, None, []

    def flush_lot_group():
        if current_rows:
            groups.append(
                {
                    "kind": "parent" if len(current_rows) > 1 else "standard",
                    "key": current_lot,
                    "rows": list(current_rows),
                }
            )

    for r in rest:
        lot = cell_str(r[C_LOT])
        # A blank/placeholder lot cannot safely group unrelated rows. Treat it
        # as a standard lot and use the RefID as the lot number at load time.
        if not lot or lot.upper() == "NEEDS LOT":
            flush_lot_group()
            current_key, current_lot, current_rows = None, None, []
            groups.append({"kind": "standard", "key": cell_str(r[C_REFID]), "rows": [r]})
            continue
        # The register occasionally reuses a lot number for a different SKU.
        # It can also reuse the same lot/SKU combination months later. Only
        # contiguous rows belong to one parent-lot group.
        key = (lot, sku_key(r))
        if current_rows and key != current_key:
            flush_lot_group()
            current_rows = []
        current_key, current_lot = key, lot
        current_rows.append(r)
    flush_lot_group()
    return groups, flags


def dedupe_register_rows(rows):
    """Keep the latest exact RefID/Lot/SKU occurrence in the import range."""
    last_index = {}
    for index, row in enumerate(rows):
        ref, lot = cell_str(row[C_REFID]), cell_str(row[C_LOT])
        if ref and lot:
            last_index[(ref.upper(), lot.upper(), sku_key(row))] = index

    kept, flags = [], []
    for index, row in enumerate(rows):
        ref, lot = cell_str(row[C_REFID]), cell_str(row[C_LOT])
        key = (ref.upper(), lot.upper(), sku_key(row))
        if ref and lot and last_index[key] != index:
            flags.append(
                f"duplicate row {row_ref(row)} RefID {ref}, Lot {lot} skipped; "
                "latest occurrence wins"
            )
            continue
        kept.append(row)
    return kept, flags


def build_test_values(row, ltt_by_name):
    out = []
    for idx, name in TEST_COLS:
        if idx < len(row):
            v = cell_str(row[idx])
            if v:
                out.append((name, ltt_by_name.get(name), v))
    return out


def count_populated_result_cells(rows):
    return sum(
        1
        for row in rows
        for idx, _name in TEST_COLS
        if idx < len(row) and cell_str(row[idx])
    )


def build_group_test_values(
    rows, ltt_by_name, flags=None, group_key=None, cell_issues=None
):
    """Merge per-test values across every source row in a lot group.

    One TestResult row is written per test type. The first non-empty source row
    wins; later same-value cells are merged and later conflicting values are
    flagged so they are visible in the load report.
    """

    out = []
    for idx, name in TEST_COLS:
        entries = []
        for row in rows:
            if idx < len(row):
                v = cell_str(row[idx])
                if v:
                    entries.append((row, v))
        if not entries:
            continue

        winner_row, winner_value = entries[0]
        ltt = ltt_by_name.get(name)
        if ltt is None and cell_issues is not None:
            for source_row, source_value in entries:
                cell_issues.append(
                    f"row {row_ref(source_row)} {name}={source_value!r} has no LabTestType mapping"
                )

        for source_row, source_value in entries[1:]:
            if source_value == winner_value:
                if cell_issues is not None:
                    cell_issues.append(
                        f"row {row_ref(source_row)} {name}={source_value!r} merged with "
                        f"row {row_ref(winner_row)}"
                    )
                continue
            msg = (
                f"conflicting {name} result in group {group_key!r}: "
                f"row {row_ref(winner_row)}={winner_value!r} wins over "
                f"row {row_ref(source_row)}={source_value!r}"
            )
            if flags is not None:
                flags.append(msg)
            if cell_issues is not None:
                cell_issues.append(
                    f"row {row_ref(source_row)} {name}={source_value!r} dropped; "
                    f"conflicts with row {row_ref(winner_row)}={winner_value!r}"
                )

        out.append((name, ltt, winner_value))
    return out


def release_date_of(group):
    for r in group["rows"]:
        d = to_date(r[C_RELEASE])
        if d:
            return d
    return None


def qc_approval_value(group, flags=None):
    values = [
        (r, cell_str(r[C_QC]))
        for r in group["rows"]
        if len(r) > C_QC and cell_str(r[C_QC])
    ]
    if not values:
        return ""
    winner_row, winner_value = values[0]
    for source_row, source_value in values[1:]:
        if source_value != winner_value and flags is not None:
            flags.append(
                f"conflicting QC Approval in group {group['key']!r}: "
                f"row {row_ref(winner_row)}={winner_value!r} used over "
                f"row {row_ref(source_row)}={source_value!r}"
            )
    return winner_value


def classify(group, ltt_by_name, test_values=None):
    """Return (LotStatus, is_released)."""
    qc = bool(qc_approval_value(group))
    needs_metals = any(
        str(r[C_TOPRINT]).strip() == "NEEDS METALS" for r in group["rows"]
    )
    has_results = bool(
        test_values
        if test_values is not None
        else build_group_test_values(group["rows"], ltt_by_name)
    )
    if qc and not needs_metals:
        return LotStatus.RELEASED, True
    if needs_metals:
        return LotStatus.PARTIAL_RESULTS, False
    if has_results:
        return LotStatus.AWAITING_RELEASE, False
    return LotStatus.AWAITING_RESULTS, False


def make_ref_generator(existing_refs=None):
    seq = defaultdict(int)
    ref_re = re.compile(r"^(\d{6})-(\d{3})$")
    for ref in existing_refs or ():
        match = ref_re.match(str(ref).strip())
        if match:
            seq[match.group(1)] = max(seq[match.group(1)], int(match.group(2)))
    today = datetime.now().date()

    def gen(mfg):
        prefix = (mfg or today).strftime("%y%m%d")
        seq[prefix] += 1
        return f"{prefix}-{seq[prefix]:03d}"

    return gen


def future_date_flags(rows, first_future_year=2029):
    flags = []
    for r in rows:
        mfg = to_date(r[C_MFG])
        if mfg and mfg.year >= first_future_year:
            flags.append(
                "future Mfg Date "
                f"{mfg.isoformat()} for RefID {cell_str(r[C_REFID])}, "
                f"Lot {cell_str(r[C_LOT])}, "
                f"{display_name(r[C_BRAND], r[C_PRODUCT], r[C_FLAVOR], r[C_SIZE])}"
            )
    return flags


def normalize_future_mfg_typos(rows):
    """Fix obvious year typos when manufacturing would occur years after release."""
    flags = []
    for r in rows:
        mfg = to_date(r[C_MFG])
        release = to_date(r[C_RELEASE])
        if not mfg or not release or mfg.year <= release.year + 1:
            continue
        try:
            corrected = mfg.replace(year=release.year)
        except ValueError:
            corrected = mfg.replace(year=release.year, day=28)
        if corrected <= release:
            r[C_MFG] = corrected
            flags.append(
                f"corrected future Mfg Date {mfg.isoformat()} -> "
                f"{corrected.isoformat()} for RefID {cell_str(r[C_REFID])}"
            )
    return flags


def append_conflicts(db, groups):
    existing_refs = {
        reference_number for (reference_number,) in db.query(Lot.reference_number).all()
    }
    existing_sublots = {
        sublot_number for (sublot_number,) in db.query(Sublot.sublot_number).all()
    }
    conflicts = []
    for g in groups:
        rows_g = g["rows"]
        first = rows_g[0]
        if g["kind"] == "standard":
            lotnum = cell_str(first[C_LOT])
            if not lotnum or lotnum.upper() == "NEEDS LOT":
                lotnum = cell_str(first[C_REFID])
            ref = cell_str(first[C_REFID])
            if ref.upper() in existing_refs:
                conflicts.append(("reference_number", ref, g["key"]))
        elif g["kind"] == "parent":
            lotnum = g.get("lotnum_override") or cell_str(first[C_LOT])
            if (
                "ref_override" in g
                and str(g["ref_override"]).strip().upper() in existing_refs
            ):
                conflicts.append(("reference_number", g["ref_override"], g["key"]))
            same_sku_cref = "ref_override" in g
            for r in rows_g:
                sublot = cell_str(r[C_LOT]) if same_sku_cref else cell_str(r[C_REFID])
                if sublot.upper() in existing_sublots:
                    conflicts.append(("sublot_number", sublot, g["key"]))
        else:
            cref = str(g["key"]).strip()
            if cref.upper() in existing_refs:
                conflicts.append(("reference_number", cref, g["key"]))
    return conflicts


def backup_sqlite_db(src_path=DB_PATH, backup_dir=BACKUP_DIR):
    if not os.path.exists(src_path):
        raise FileNotFoundError(f"database not found: {src_path}")
    os.makedirs(backup_dir, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = os.path.join(backup_dir, f"labtrack.db.bak-{ts}")
    with sqlite3.connect(src_path) as src_conn, sqlite3.connect(backup) as dest_conn:
        src_conn.backup(dest_conn)
    if not os.path.exists(backup):
        raise RuntimeError("backup not created")
    return backup


def wipe_per_lot_data(db):
    for table_name in WIPE_ORDER:
        db.execute(text(f"DELETE FROM {table_name}"))


def print_reconciliation(db, rows, counts, populated_cells, cell_issues):
    db.expire_all()
    lot_count = db.query(Lot).count()
    sublot_count = db.query(Sublot).count()
    result_count = db.query(TestResult).count()
    release_count = db.query(COARelease).count()
    rows_with_qc = sum(1 for row in rows if len(row) > C_QC and not is_blank(row[C_QC]))

    print("\nReconciliation:")
    print(f"  Spreadsheet rows seen        : {len(rows)}")
    print(
        f"  Lots + sublots in DB         : {lot_count + sublot_count} ({lot_count} lots, {sublot_count} sublots)"
    )
    print(f"  Populated result cells       : {populated_cells}")
    print(f"  TestResult rows in DB        : {result_count}")
    print(f"  Rows with QC Approval        : {rows_with_qc}")
    print(f"  COARelease rows in DB        : {release_count}")
    print(f"  Products auto-created        : {counts['products']}")

    if cell_issues:
        print("  Result cells merged/dropped/unmapped:")
        for issue in cell_issues:
            print(f"    ! {issue}")

    warnings = []
    if len(rows) != lot_count + sublot_count:
        warnings.append(
            f"spreadsheet rows ({len(rows)}) != lots+sublots ({lot_count + sublot_count})"
        )
    if populated_cells != result_count:
        warnings.append(
            f"populated result cells ({populated_cells}) != TestResult rows ({result_count})"
        )
    if rows_with_qc != release_count:
        warnings.append(
            f"rows with QC Approval ({rows_with_qc}) != COARelease rows ({release_count})"
        )
    if any("dropped" in issue or "unmapped" in issue for issue in cell_issues):
        warnings.append("some populated result cells were dropped or unmapped")

    if warnings:
        print("\n!!! WARNING: RECONCILIATION MISMATCH !!!")
        for warning in warnings:
            print(f"!!! {warning}")
        print("!!! Review the load report before go-live.")


# --------------------------------------------------------------------------- #
def run_load(file_path=DEFAULT_XLSX, commit=False, dry_run=False, append_from_row=None):
    commit = commit and not dry_run
    append_mode = append_from_row is not None

    print(f"Source : {file_path}")
    print(f"DB     : {DB_PATH}")
    if append_mode:
        mode = "APPEND COMMIT" if commit else "APPEND DRY RUN"
        print(f"Mode   : {mode} (from Excel row {append_from_row}; no wipe)\n")
    else:
        print(
            f"Mode   : {'COMMIT (wipe + load)' if commit else 'DRY RUN (no writes)'}\n"
        )

    rows = load_rows(file_path, append_from_row)
    print(f"Rows read: {len(rows)}")
    rows, flags = dedupe_register_rows(rows)
    groups, group_flags = group_rows(rows)
    flags.extend(group_flags)
    flags.extend(normalize_future_mfg_typos(rows))
    flags.extend(future_date_flags(rows))
    populated_cells = count_populated_result_cells(rows)

    # per-product test panel (union across the product's rows) + needs-metals marker
    product_tests, product_needs_metals = defaultdict(set), defaultdict(bool)
    for r in rows:
        k = sku_key(r)
        for idx, name in TEST_COLS:
            if idx < len(r) and cell_str(r[idx]):
                product_tests[k].add(name)
        if str(r[C_TOPRINT]).strip() == "NEEDS METALS":
            product_needs_metals[k] = True

    db = SessionLocal()
    try:
        ltt_by_name = {t.test_name: t for t in db.query(LabTestType).all()}
        qc_user = db.query(User).filter(User.username == "qcmanager").first()
        if qc_user is None:
            print("FATAL: qcmanager user not found")
            sys.exit(1)
        existing = {
            (
                (p.brand or "").strip().lower(),
                (p.product_name or "").strip().lower(),
                (p.flavor or "").strip().lower(),
                (p.size or "").strip().lower(),
            ): p.id
            for p in db.query(Product).all()
        }

        statuses = Counter()
        cell_issues = []
        group_test_values = []
        for g in groups:
            values = build_group_test_values(
                g["rows"],
                ltt_by_name,
                flags=flags,
                group_key=g["key"],
                cell_issues=cell_issues,
            )
            group_test_values.append(values)
            qc_approval_value(g, flags=flags)
            st, _ = classify(g, ltt_by_name, values)
            statuses[st.value] += 1
        needed = OrderedDict()
        for g in groups:
            for r in g["rows"]:
                needed.setdefault(sku_key(r), r)
        to_create = [r for k, r in needed.items() if k not in existing]

        n = {
            kind: sum(1 for g in groups if g["kind"] == kind)
            for kind in ("standard", "parent", "composite")
        }
        print(
            f"\nLots: {len(groups)}  (standard {n['standard']}, parent {n['parent']}, composite {n['composite']})"
        )
        print(
            "Status breakdown: "
            + ", ".join(f"{k}={v}" for k, v in sorted(statuses.items()))
        )
        print(
            f"Products: {len(needed)} SKUs referenced, {len(needed) - len(to_create)} existing, {len(to_create)} to create"
        )

        conflicts = append_conflicts(db, groups) if append_mode else []
        if conflicts:
            print("\nFATAL: append import would collide with existing identifiers:")
            for field, value, group_key in conflicts:
                print(f"  ! {field} {value!r} in group {group_key!r}")
            sys.exit(1)

        if not commit:
            if flags:
                print("\nFlags:")
                for f in flags:
                    print(f"  ! {f}")
            if cell_issues:
                print("\nResult cell merge/drop/unmapped report:")
                for issue in cell_issues:
                    print(f"  ! {issue}")
            print("\nDRY RUN complete — no writes. Re-run with --commit to apply.")
            return

        backup = backup_sqlite_db(DB_PATH, BACKUP_DIR)
        print(f"\nBacked up DB -> {backup}")
        if append_mode:
            print("Append mode: existing per-lot lab data left intact.")
        else:
            wipe_per_lot_data(db)
            print("Wiped per-lot lab data.")

        existing_refs = (
            {
                reference_number
                for (reference_number,) in db.query(Lot.reference_number).all()
            }
            if append_mode
            else set()
        )
        existing_sublots = (
            {sublot_number for (sublot_number,) in db.query(Sublot.sublot_number).all()}
            if append_mode
            else set()
        )
        existing_lot_numbers = (
            {lot_number for (lot_number,) in db.query(Lot.lot_number).all()}
            if append_mode
            else set()
        )
        gen_ref = make_ref_generator(existing_refs)
        used_lot_numbers = set(existing_lot_numbers)
        used_refs, used_sublots = set(existing_refs), set(existing_sublots)
        pid_cache = dict(existing)
        counts = Counter()

        specs_done = set()

        def ensure_specs(pid, k):
            # Give a product its required panel, but never touch a product that
            # already has specs (e.g. the seed catalog). Covers products created
            # by an earlier run that had no specs yet.
            if pid in specs_done:
                return
            specs_done.add(pid)
            if db.query(ProductTestSpecification).filter_by(product_id=pid).count() > 0:
                return
            panel = set(product_tests.get(k) or ()) or set(MICRO_COMMON)
            if product_needs_metals.get(k):
                panel |= set(METALS)
            for name in panel:
                ltt = ltt_by_name.get(name)
                if ltt:
                    db.add(
                        ProductTestSpecification(
                            product_id=pid,
                            lab_test_type_id=ltt.id,
                            is_required=True,
                            specification=(
                                ltt.default_specification or "Within limits"
                            )[:100],
                        )
                    )
                    counts["specs"] += 1

        def get_pid(row):
            k = sku_key(row)
            if k in pid_cache:
                ensure_specs(pid_cache[k], k)
                return pid_cache[k]
            p = Product(
                brand=str(row[C_BRAND]).strip(),
                product_name=str(row[C_PRODUCT]).strip(),
                flavor=(
                    (str(row[C_FLAVOR]).strip() or None)
                    if not is_blank(row[C_FLAVOR])
                    else None
                ),
                size=(
                    (str(row[C_SIZE]).strip() or None)
                    if not is_blank(row[C_SIZE])
                    else None
                ),
                display_name=display_name(
                    row[C_BRAND], row[C_PRODUCT], row[C_FLAVOR], row[C_SIZE]
                ),
            )
            db.add(p)
            db.flush()
            pid_cache[k] = p.id
            counts["products"] += 1
            ensure_specs(p.id, k)
            return p.id

        def uniq(value, used, label):
            v, i = value, 1
            while v in used:
                i += 1
                v = f"{value}-{i}"
            if v != value:
                flags.append(f"deduped {label} {value!r} -> {v!r}")
            used.add(v)
            return v

        def add_results(lot_id, values, approved, rel_date):
            for name, ltt, v in values:
                tr = TestResult(
                    lot_id=lot_id,
                    test_type=name,
                    result_value=v,
                    unit=(ltt.default_unit if ltt else None),
                    specification=(ltt.default_specification if ltt else None),
                    method=(ltt.test_method if ltt else None),
                    lab_test_type_id=(ltt.id if ltt else None),
                    include_on_coa=True,
                    status=(
                        TestResultStatus.APPROVED
                        if approved
                        else TestResultStatus.DRAFT
                    ),
                )
                if approved:
                    tr.approved_by_id = qc_user.id
                    tr.approved_at = datetime.combine(
                        rel_date or datetime.now().date(), datetime.min.time()
                    )
                db.add(tr)
                counts["results"] += 1

        def add_release(lot_id, pid, rel_date, qc_name):
            note = "Loaded from COA register"
            if qc_name:
                note = f"{note} (QC: {qc_name})"
            db.add(
                COARelease(
                    lot_id=lot_id,
                    product_id=pid,
                    status=COAReleaseStatus.RELEASED,
                    released_at=datetime.combine(
                        rel_date or datetime.now().date(), datetime.min.time()
                    ),
                    released_by_id=qc_user.id,
                    notes=note,
                )
            )
            counts["releases"] += 1

        for g, test_values in zip(groups, group_test_values):
            rows_g = g["rows"]
            first = rows_g[0]
            status, _ = classify(g, ltt_by_name, test_values)
            approved = status in (LotStatus.RELEASED, LotStatus.AWAITING_RELEASE)
            rel_date = release_date_of(g)
            qc_name = qc_approval_value(g)

            if g["kind"] == "standard":
                lotnum = cell_str(first[C_LOT])
                if not lotnum or lotnum.upper() == "NEEDS LOT":
                    lotnum = cell_str(first[C_REFID])
                    flags.append(f"blank/NEEDS-LOT row: lot_number set to RefID {lotnum}")
                lotnum = uniq(lotnum, used_lot_numbers, "lot number")
                source_ref = cell_str(first[C_REFID])
                if not source_ref:
                    source_ref = gen_ref(to_date(first[C_MFG]))
                    flags.append(
                        f"blank RefID row {row_ref(first)}: generated reference {source_ref}"
                    )
                ref = uniq(source_ref, used_refs, "reference")
                lot = Lot(
                    lot_number=lotnum,
                    lot_type=LotType.STANDARD,
                    reference_number=ref,
                    mfg_date=to_date(first[C_MFG]),
                    exp_date=to_date(first[C_EXP]),
                    status=status,
                    generate_coa=True,
                )
                db.add(lot)
                db.flush()
                db.add(LotProduct(lot_id=lot.id, product_id=get_pid(first)))
                add_results(lot.id, test_values, approved, rel_date)
                if status == LotStatus.RELEASED:
                    add_release(lot.id, get_pid(first), rel_date, qc_name)

            elif g["kind"] == "parent":
                lotnum = g.get("lotnum_override") or cell_str(first[C_LOT])
                lotnum = uniq(lotnum, used_lot_numbers, "lot number")
                mfgs = [to_date(r[C_MFG]) for r in rows_g if to_date(r[C_MFG])]
                mfg = min(mfgs) if mfgs else None
                ref = (
                    uniq(g["ref_override"], used_refs, "reference")
                    if "ref_override" in g
                    else uniq(gen_ref(mfg), used_refs, "reference")
                )
                lot = Lot(
                    lot_number=lotnum,
                    lot_type=LotType.PARENT_LOT,
                    reference_number=ref,
                    mfg_date=mfg,
                    exp_date=to_date(first[C_EXP]),
                    status=status,
                    generate_coa=True,
                )
                db.add(lot)
                db.flush()
                db.add(LotProduct(lot_id=lot.id, product_id=get_pid(first)))
                same_sku_cref = "ref_override" in g
                for r in rows_g:
                    src = (
                        str(r[C_LOT]).strip()
                        if same_sku_cref
                        else str(r[C_REFID]).strip()
                    )
                    db.add(
                        Sublot(
                            parent_lot_id=lot.id,
                            sublot_number=uniq(src, used_sublots, "sublot"),
                            production_date=to_date(r[C_MFG]),
                        )
                    )
                    counts["sublots"] += 1
                add_results(lot.id, test_values, approved, rel_date)
                if status == LotStatus.RELEASED:
                    add_release(lot.id, get_pid(first), rel_date, qc_name)

            else:  # composite
                cref = g["key"]
                lotnum = uniq(cref, used_lot_numbers, "lot number")
                ref = uniq(cref, used_refs, "reference")
                lot = Lot(
                    lot_number=lotnum,
                    lot_type=LotType.MULTI_SKU_COMPOSITE,
                    reference_number=ref,
                    mfg_date=to_date(first[C_MFG]),
                    exp_date=to_date(first[C_EXP]),
                    status=status,
                    generate_coa=True,
                )
                db.add(lot)
                db.flush()
                by_pid = OrderedDict()
                for r in rows_g:
                    by_pid.setdefault(get_pid(r), []).append(str(r[C_LOT]).strip())
                for pid, lots in by_pid.items():
                    batch_number = ", ".join(lots)
                    if any(len(lot_number) > 50 for lot_number in lots):
                        flags.append(
                            f"component batch number exceeds 50 chars in composite {cref}: {lots!r}"
                        )
                    if len(batch_number) > 50:
                        flags.append(
                            f"composite {cref} batch_number exceeds 50 chars for product_id {pid}: "
                            f"{batch_number!r}"
                        )
                    db.add(
                        LotProduct(
                            lot_id=lot.id, product_id=pid, batch_number=batch_number
                        )
                    )
                add_results(lot.id, test_values, approved, rel_date)
                if status == LotStatus.RELEASED:
                    for pid in by_pid:
                        add_release(lot.id, pid, rel_date, qc_name)

        db.commit()
        print(
            f"\nLoaded: {len(groups)} lots, {counts['sublots']} sublots, {counts['results']} test results, "
            f"{counts['releases']} releases, {counts['products']} products created, {counts['specs']} specs added."
        )
        print(
            "Status breakdown: "
            + ", ".join(f"{k}={v}" for k, v in sorted(statuses.items()))
        )
        print_reconciliation(db, rows, counts, populated_cells, cell_issues)
        if flags:
            print("\nFlags:")
            for f in flags:
                print(f"  ! {f}")
        if counts["releases"]:
            print(
                "\nPost-load reminder: released COAs need immutable snapshots. "
                "Run scripts/backfill_coa_snapshots.py --commit before go-live use."
            )
        print(f"\nDone. Backup at {backup}")
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def main():
    ap = argparse.ArgumentParser(description="Wipe and load the COA register")
    ap.add_argument("--file", default=DEFAULT_XLSX)
    ap.add_argument("--commit", action="store_true")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument(
        "--append-from-row",
        type=int,
        help="Append only nonblank workbook rows at or after this Excel row number; never wipes existing lots.",
    )
    args = ap.parse_args()
    run_load(
        file_path=args.file,
        commit=args.commit,
        dry_run=args.dry_run,
        append_from_row=args.append_from_row,
    )


if __name__ == "__main__":
    main()
